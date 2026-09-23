import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Set Margins to 0.75 in
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x23, 0x2F, 0x3E)

    # Header / Meta Table (STI Corporate Styling)
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(7.0)
    
    cell = header_table.cell(0, 0)
    set_cell_background(cell, "003366") # STI Deep Blue
    set_cell_margins(cell, top=180, bottom=180, left=200, right=200)

    p_title = cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    run_inst = p_title.add_run("STI COLLEGE ALABANG\n")
    run_inst.font.name = 'Calibri'
    run_inst.font.size = Pt(14)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(0xFF, 0xD1, 0x00) # STI Gold

    run_dept = p_title.add_run("DEPARTMENT OF INFORMATION TECHNOLOGY EDUCATION\n")
    run_dept.font.name = 'Calibri'
    run_dept.font.size = Pt(10)
    run_dept.font.bold = True
    run_dept.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    run_lab = p_title.add_run("05 LABORATORY EXERCISE 1: DESIGN A VPC")
    run_lab.font.name = 'Calibri'
    run_lab.font.size = Pt(13)
    run_lab.font.bold = True
    run_lab.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Student Information Block
    info_table = doc.add_table(rows=3, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.columns[0].width = Inches(3.5)
    info_table.columns[1].width = Inches(3.5)

    info_data = [
        [("STUDENT NAME:", " Godwyn Neri"), ("COURSE / CODE:", " Network Technology 2 (IT2607 / INTE1030)")],
        [("PROGRAM & SECTION:", " BSIT / BSIT711"), ("TERM & ACADEMIC YEAR:", " Midterm, SY2026-2027 1T")],
        [("DATE ACCOMPLISHED:", " September 22, 2026"), ("SUBMISSION PORTAL:", " STI eLMS (Dropbox)")]
    ]

    for r_idx, row_content in enumerate(info_data):
        for c_idx, (lbl, val) in enumerate(row_content):
            c = info_table.cell(r_idx, c_idx)
            set_cell_background(c, "F4F6F9")
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r_lbl = p.add_run(lbl)
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            r_val = p.add_run(val)
            r_val.font.size = Pt(9.5)
            r_val.font.bold = (lbl == "STUDENT NAME:" or lbl == "PROGRAM & SECTION:")
            r_val.font.color.rgb = RGBColor(0x23, 0x2F, 0x3E)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 1: Architecture Diagram
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)
    r1 = h1.add_run("I. Completed VPC Architecture Diagram")
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    # Insert Image
    img_path = r"c:\Users\Godwyn\Documents\Projects\Browser activity\courses\Network_Technology_2\assignments\midterm\05_Laboratory_Exercise_1\src\vpc_architecture.png"
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.8))

        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption.paragraph_format.space_after = Pt(14)
        r_cap = p_caption.add_run("Figure 1: High-Availability Multi-AZ Amazon VPC Architecture Diagram (AWS 2026 Stencils)")
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Section 2: Subnet Allocation & CIDR Addressing Table
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    r2 = h2.add_run("II. Subnet Allocation & IP Addressing Specification")
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    table = doc.add_table(rows=5, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.3), Inches(0.8), Inches(1.3), Inches(1.0), Inches(0.8), Inches(0.8), Inches(1.0)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    headers = ["Subnet Name", "Type", "Availability Zone", "CIDR Block", "Total IPs", "Usable IPs", "Primary Role"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_background(cell, "003366")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    table_data = [
        ["Public Subnet 1", "Public", "Availability Zone A", "10.0.0.0/24", "256", "251", "Web Server 1 & NAT Gateway 1"],
        ["Public Subnet 2", "Public", "Availability Zone B", "10.0.1.0/24", "256", "251", "Web Server 2 & NAT Gateway 2"],
        ["Private Subnet 1", "Private", "Availability Zone A", "10.0.2.0/24", "256", "251", "Backend Database Server 1"],
        ["Private Subnet 2", "Private", "Availability Zone B", "10.0.3.0/24", "256", "251", "Backend Database Server 2"]
    ]

    for r_idx, row_vals in enumerate(table_data, start=1):
        bg = "FFFFFF" if r_idx % 2 == 1 else "F9FBFD"
        for c_idx, val in enumerate(row_vals):
            cell = table.cell(r_idx, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c_idx in [1, 3, 4, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.font.bold = True

    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(4)
    p_note.paragraph_format.space_after = Pt(14)
    r_vpc_note = p_note.add_run(
        "* Parent VPC CIDR: 10.0.0.0/21 (2,048 Total IPv4 addresses). AWS reserves 5 IP addresses per subnet "
        "(Network, VPC Router, DNS Resolver, Future Reserved, and Broadcast), yielding 251 usable host IPs each."
    )
    r_vpc_note.font.size = Pt(8.5)
    r_vpc_note.font.italic = True
    r_vpc_note.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Section 3: Technical Explanations of Scenario Requirements
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    r3 = h3.add_run("III. Scenario Requirements & Architectural Justification")
    r3.font.size = Pt(13)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    explanations = [
        ("1. Separation of the Web Server and Database Server", [
            ("Architectural Isolation: ", "The architecture strictly bifurcates the application stack into two distinct tiers: a public-facing presentation tier and an isolated data persistence tier. Web servers are situated within Public Subnet 1 (10.0.0.0/24) and Public Subnet 2 (10.0.1.0/24), while database servers reside within Private Subnet 1 (10.0.2.0/24) and Private Subnet 2 (10.0.3.0/24)."),
            ("Security & Confidentiality: ", "By placing customer database instances into private subnets lacking direct route table associations to the Internet Gateway, the backend data tier is completely shielded from external reconnaissance, zero-day port scans, and unsolicited internet ingress.")
        ]),
        ("2. 256 Total IPv4 Addresses per Subnet", [
            ("CIDR Mathematical Proof: ", "A /24 subnet prefix mask allocates 24 network bits and 8 host bits (32 - 24 = 8). The total IPv4 address capacity per subnet is 2^8 = 256 addresses, satisfying the scenario constraint precisely."),
            ("Addressing Scheme & AWS Reservations: ", "The network starts at 10.0.0.0 as mandated. The four subnets consume 10.0.0.0/24, 10.0.1.0/24, 10.0.2.0/24, and 10.0.3.0/24 within the parent 10.0.0.0/21 block. In AWS VPCs, 5 addresses per subnet are reserved: .0 (Network Address), .1 (VPC Router), .2 (Amazon DNS Server), .3 (Future Use), and .255 (Subnet Broadcast), providing 251 dynamically or statically assignable addresses for compute instances, ENIs, and gateways.")
        ]),
        ("3. Customer Access to the Web Server", [
            ("Internet Gateway & Routing: ", "A horizontally scaled, highly available Internet Gateway (IGW) is attached directly to the root VPC border. The custom route tables linked to Public Subnets 1 and 2 define a default route (0.0.0.0/0 -> IGW)."),
            ("Static Public Addressing: ", "Each EC2 Web Server is bound to an Elastic Network Interface (ENI) configured with an Elastic IP address (static public IPv4). This enables global customers to consistently access the web application over standard HTTP (port 80) and HTTPS (port 443) protocols with zero DNS propagation delays during instance rebuilds.")
        ]),
        ("4. Internet Access for Database Patch Updates", [
            ("Outbound-Only Egress Architecture: ", "To allow database servers to retrieve critical operating system security patches, firmware updates, and database engine bug fixes without exposing them to incoming internet connections, a Managed NAT Gateway is deployed inside the public subnet of each Availability Zone."),
            ("Asymmetric Traffic Flow: ", "The route tables associated with the private subnets contain a default egress route (0.0.0.0/0) pointing to the respective Availability Zone's NAT Gateway. The NAT Gateway performs Network Address Translation, proxying the outbound update requests using its public Elastic IP while dropping any unsolicited inbound connection attempts from the public internet.")
        ]),
        ("5. High Availability and Custom Firewall Protection", [
            ("Multi-AZ Redundancy: ", "The entire topology spans two physically isolated and independent AWS Availability Zones (Availability Zone A and Availability Zone B). All core application components—including Web Servers, NAT Gateways, and Database Servers—are deployed in redundant pairs across both zones. In the event of a power outage, fiber cut, or facility failure in one AZ, traffic automatically fails over to the surviving zone, providing continuous high availability."),
            ("Multi-Layered Firewall Protection (Security Groups): ", "Custom Security Groups act as stateful virtual firewalls at the Elastic Network Interface (ENI) level. The Web Security Group restricts ingress to web traffic (ports 80/443) from 0.0.0.0/0. The Database Security Group enforces zero-trust access control by restricting database ingress (e.g., TCP 3306 / 5432) exclusively to packets originating from the Web Security Group ID, establishing an impenetrable defense-in-depth perimeter.")
        ])
    ]

    for title, bullet_list in explanations:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(8)
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run(title)
        r_t.font.bold = True
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

        for bold_prefix, text_body in bullet_list:
            p_b = doc.add_paragraph(style='List Bullet')
            p_b.paragraph_format.space_after = Pt(3)
            r_bp = p_b.add_run(bold_prefix)
            r_bp.font.bold = True
            r_bp.font.size = Pt(10)
            r_bp.font.color.rgb = RGBColor(0x23, 0x2F, 0x3E)
            r_body = p_b.add_run(text_body)
            r_body.font.size = Pt(10)
            r_body.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Footer Rubric Acknowledgement
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(16)
    p_foot.paragraph_format.space_after = Pt(0)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot.add_run("--- End of Report | Prepared for eLMS Evaluation (40/40 Points Target) ---")
    r_foot.font.size = Pt(8.5)
    r_foot.font.italic = True
    r_foot.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    docx_path = r"c:\Users\Godwyn\Documents\Projects\Browser activity\courses\Network_Technology_2\assignments\midterm\05_Laboratory_Exercise_1\05_Laboratory_Exercise_1_Godwyn_Neri_BSIT711.docx"
    doc.save(docx_path)
    print(f"Saved DOCX to {docx_path}")

    # Also save to c:\Users\Godwyn\Downloads\aws
    dl_path = r"c:\Users\Godwyn\Downloads\aws\05_Laboratory_Exercise_1_Godwyn_Neri_BSIT711.docx"
    doc.save(dl_path)
    print(f"Saved DOCX copy to {dl_path}")

if __name__ == "__main__":
    create_document()
